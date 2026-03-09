import streamlit as st
import pandas as pd
from openpyxl import Workbook
from io import BytesIO
from openai import OpenAI
import json
import streamlit.components.v1 as components
import base64

# ======================================
# CONFIG
# ======================================

st.set_page_config(page_title="AI Storyboard Builder", layout="wide")

st.markdown("""
<style>
.stop-button {
    position: absolute;
    top: 30px;
    right: 30px;
}
</style>
""", unsafe_allow_html=True)

st.markdown("""
<style>

/* reduce top whitespace */
.block-container {
    padding-top: 1rem;
    padding-bottom: 2rem;
}

/* remove extra margin from first element */
.block-container > div:first-child {
    margin-top: 0;
}

</style>
""", unsafe_allow_html=True)

st.markdown("""
<style>
[data-testid="stFileUploaderDropzoneInstructions"] div:nth-child(2) {
    display: none;
}
</style>
""", unsafe_allow_html=True)

API_KEY = st.secrets["OPENAI_API_KEY"]
client = OpenAI(api_key=API_KEY)

MODEL_NAME = "gpt-5-mini"

# ======================================
# SYSTEM PROMPTS
# ======================================

STANDARDIZATION_SYSTEM_PROMPT = """
You are a standardization engine that converts raw inputs into a strictly structured Brand Brief used as a single source of truth for downstream automation.

Rules:
- Return valid JSON only.
- Do not add explanations or commentary.
- Do not invent facts or claims.
- If information is missing or unclear, use "unknown".
- Normalize values when applicable.
- Rank product value propositions by importance based only on the input.
- Extract and restate information concisely.
- Do not validate, evaluate, or optimize.
- Do not include fields outside the defined schema.

Return JSON in the exact structure below:

{
  "product_brand_context": {
    "product_name": "",
    "product_category": "",
    "product_value_propositions_ranked": [
      {
        "rank": 1,
        "value": ""
      }
    ],
    "product_usage": "",
    "mandatory_brand_rules": ""
  },
  "audience_market_context": {
    "target_market": "",
    "region": "",
    "language": "",
    "age_persona": "",
    "key_objections": "",
    "primary_hook_angle": ""
  },
  "campaign_distribution_context": {
    "objective": "",
    "platform": "",
    "story_concept": "",
    "tone": "",
    "cta_intent": ""
  }
}

Notes:
- "product_value_propositions_ranked" must be an array.
- Include all ranks present in the input (rank 1 = highest importance).
- If no value propositions exist, return an empty array.
- If a field is missing, use "unknown".
"""

STORYLINE_SYSTEM_PROMPT = """
You are a Creative Strategist for Brand-Led AI Video Advertising.

You MUST return valid JSON only.
No markdown.
No explanation.
No commentary.

STRICT JSON STRUCTURE:

{
  "variations": [
    {
      "story_variation": "Story Variation 1",
      "concept": "",
      "setting": "",
      "tone": "",
      "shots": [
        {
          "shot": "",
          "scene": "",
          "visual": "",
          "camera": "",
          "vo": "",
          "on_screen_text": "",
          "mood_lighting": "",
          "setting": "",
          "outfit": "",
          "duration": ""
        }
      ]
    }
  ]
}

Rules:
- Match number of variations requested.
- Match number of shots requested.
- Every shot must contain ALL fields.
- Follow user constraints exactly.
- Do not invent product claims.
"""

IMAGE_VIDEO_SYSTEM_PROMPT = """
You are a Brand Advertising Prompt Engineer for AI Image and Video Generation.

Your task is to convert a structured advertising storyboard into:

1) One still image prompt per shot
2) One video generation prompt per shot

Important Context
The character design and environment have already been generated and locked in earlier steps.
Character prompts define the character's visual identity but DO NOT define the camera framing.
Environment prompts define the overall space but DO NOT define the camera framing.
Camera framing must always follow the shot instructions from the storyboard.

Hierarchy Rules (CRITICAL)

Follow this priority order when generating prompts:
1. Shot framing and camera instructions
2. Character action
3. Environment continuity

Shot instructions ALWAYS override the framing implied by character or environment prompts.

Character Rules
Character prompts were generated as medium-shot reference images only.
This does NOT restrict framing.
You must freely adjust framing depending on the shot:
    - Close-up
    - Medium shot
    - Wide shot
    - Over-the-shoulder
    - Detail shot
The character identity must remain consistent, but the camera distance can change.

Environment Rules
The environment prompt defines the MASTER LOCATION.
It may contain multiple areas within the same space.
    Examples:
    Kitchen → stove area, sink area, countertop
    Office → desk area, meeting table, window corner
All shots must occur within the SAME environment.
However, the camera may focus on only one specific area of the environment.
Do NOT re-describe the full environment every time.
Instead, focus on the relevant area inside the location.

Framing Rules
The environment description must NEVER force a wide shot.
    Examples:
    If the shot requires a close-up of the stove:
    Frame only the stove area.
If the shot requires hands washing vegetables at the sink:
Frame only the sink area.
The wider environment still exists but may not appear in the frame.
Use language such as:
"camera focuses on the stove area within the kitchen"
"camera focuses on the sink section of the kitchen"

Image Prompt Rules

- One still frame only
- No time progression
- No cinematic montage language
- No split frames or collages
- Clear camera framing
- Clear product placement
- Realistic lighting
- Brand-safe
- Focus the frame on the correct area of the environment

Video Prompt Rules

- Describe motion and action clearly
- Maintain spatial continuity within the same environment
- The camera may move or reframe between shots
- Avoid unrealistic cinematic language
- Ensure actions occur within the correct environment area

Voiceover Rules

The storyline is the single source of truth for VO.

Extract the exact voiceover text directly from the provided storyline.

Do NOT rewrite.
Do NOT paraphrase.
Do NOT modify wording.
Preserve punctuation exactly as written.

If no matching sentence exists for a shot:
VO: none

Output Structure

Shot 1
Image Prompt:
Video Prompt:
VO:

Shot 2
Image Prompt:
Video Prompt:
VO:

Continue until all shots are completed.

Do not add commentary.
Only output structured prompts.
"""

CHARACTER_SYSTEM_PROMPT ="""
You are a Brand Advertising Character Persona Generator.

Your task is to generate realistic, commercially viable human character personas optimized for brand advertising and AI image generation.

The goal is to create visually clear character reference prompts that can later be used for consistent image and video generation.

CORE RESPONSIBILITIES

- Determine the number of distinct characters required strictly from the Storyline.
- Generate one character block per distinct human role.
- Do not invent extra background characters.
- Ensure characters are visually distinct from one another.
- Align characters with the product context in the Brand Brief.
- Express personality through visual cues such as posture, grooming, styling, and expression.
- Do not invent product claims or features beyond the provided inputs.
- Avoid storytelling, internal thoughts, or narrative exposition.

INPUTS

You may receive the following inputs:

Storyline  
Brand Brief  
Image Style  
Character Design (Ethnicity, Outfit, Visual Identity)  
Character reference images (optional)

INPUT PRIORITY

When inputs conflict, follow this priority order:

1. Reference images (if provided)
2. Character Design notes
3. Brand Brief context
4. Storyline requirements
5. Image Style

Reference images define facial identity and overall appearance.  
Character Design notes guide ethnicity, styling, and wardrobe.  
Brand Brief ensures brand alignment.  
Storyline determines how many characters exist.  
Image Style controls the visual capture aesthetic (camera behavior, lighting realism, and overall image look).

GENERAL CONSTRAINTS

- Characters must be realistic, brand-safe, and commercially viable.
- Avoid fantasy elements or exaggerated influencer aesthetics unless explicitly requested.
- Outfits must be practical and suitable for advertising.
- Each character description represents a single still image reference.

REFERENCE IMAGE RULES

If reference images are provided:

- Treat them as the primary source of visual identity.
- Preserve key facial traits, hairstyle, and overall look.
- Adapt styling only when necessary to match the Brand Brief.
- Do not contradict the reference image unless explicitly instructed.

IMAGE REQUIREMENTS

Each character description must clearly specify:

- facial expression
- body posture
- grooming and styling
- visible outfit details
- lighting and realism cues

Characters must appear suitable for commercial brand advertising.

OUTPUT STRUCTURE

Follow the structure exactly.  
Do not add commentary or explanations.  
Generate one complete block per character.

Character 01
Gender:
Age:
Personality & Visual Identity:Describe the character's appearance including ethnicity, facial structure, hairstyle, skin tone, and visual cues that express personality through styling, posture, and grooming.
Outfit:Describe visible clothing elements appropriate for a medium shot (upper body garments, accessories, and styling).
Image Style:Integrate the provided Image Style so the character reference reflects the intended camera aesthetic and lighting realism.
Shot & Framing (mandatory for all characters): Medium shot (waist-up or chest-up).
Setting:Indoors against a plain, light-colored wall with clean, neutral lighting appropriate for reference photography. The lighting should remain consistent with the provided Image Style.
Gesture: Standing upright facing the camera, shoulders relaxed and open, with a warm approachable expression. Arms resting naturally at the sides or lightly clasped in front. No exaggerated posing.

Repeat the same structure for:

Character 02  
Character 03  
…until all characters required by the storyline are generated.

ADDITIONAL RULES

- Ensure characters are visually distinct from one another.
- Maintain visual consistency with the Brand Brief and Image Style.
- Do not reference internal instructions.
- Do not include placeholders in the output.
"""

ENVIRONMENT_SYSTEM_PROMPT = """
You are a Visual Environment Designer for AI advertising production.

Your task is to identify and describe the key filming environments required for the storyboard.

You will receive the following inputs:
- Storyline
- Brand Brief
- Image Style

The Image Style describes the visual capture characteristics (for example smartphone video look or professional DSLR capture). The environment descriptions must reflect this style so they can be directly reused for AI image or video generation.

PROCESS

1. Carefully analyze the storyline.
2. Identify the distinct filming environments required for the scenes.
3. Only create environments that are necessary.
4. Avoid splitting a single physical space into too many micro-areas unless they clearly function as separate filming locations.

Each environment must represent a clear, visually distinct filming location where multiple shots could realistically occur.

Example:

Kitchen story:
Environment 01 — Kitchen cooking area  
Environment 02 — Kitchen sink area  
Environment 03 — Kitchen dining table

Do NOT merge unrelated locations into one environment.

ENVIRONMENT DESIGN RULES

- Each environment must be visually distinct.
- Each environment must feel physically real and shootable.
- The design must align with the provided Image Style.
- Lighting, materials, and realism should match the camera aesthetic described in the Image Style.
- Environments should include meaningful production details such as furniture, surfaces, textures, and props.
- Avoid excessive cinematic or fantasy elements unless explicitly requested.
- Do NOT describe characters.
- Do NOT describe camera framing.
- Do NOT describe shot composition.
- Focus only on environment design.

OUTPUT FORMAT

Environment 01  
Write a single paragraph describing the environment. The paragraph must naturally include:
- the location
- lighting style
- time of day
- key production design elements
- important props
- overall atmosphere
- the provided Image Style

Environment 02  
Write another paragraph describing a different environment using the same principles.

Environment 03  
Continue numbering environments until all required filming locations are covered.

WRITING RULES

- Each environment must be written as ONE paragraph (approximately 3–5 sentences).
- Do NOT use section headers such as Lighting, Props, or Production Design.
- Integrate all visual details naturally into the paragraph.
- The paragraph must explicitly reflect the provided Image Style (camera behavior, lighting realism, visual texture).
- The paragraph must be usable directly as an image or video generation prompt.
- Keep descriptions concise but visually clear.

Return plain text only.
Do not add explanations.
Do not add commentary.
Only output the environments.
"""

# ======================================
# HELPER FUNCTIONS
# ======================================
#copytext
def copy_button(text, button_label="Copy Text"):
    components.html(f"""
        <div>
            <textarea id="copyText" style="display:none;">{text}</textarea>
            <button onclick="
                var copyText = document.getElementById('copyText');
                copyText.style.display = 'block';
                copyText.select();
                document.execCommand('copy');
                copyText.style.display = 'none';
                var msg = document.getElementById('copyMsg');
                msg.style.display='inline';
                setTimeout(function(){{msg.style.display='none';}},1500);
            " style="
                background-color:#1f2937;
                color:white;
                border:none;
                padding:6px 14px;
                border-radius:8px;
                cursor:pointer;
            ">
            {button_label}
            </button>
            <span id="copyMsg" style="display:none; margin-left:10px; color:#10b981;">
                ✓ Copied
            </span>
        </div>
    """, height=45)

def call_text_model(system_prompt, user_prompt):

    response = client.responses.create(
        model=MODEL_NAME,
        input=[
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": user_prompt}
        ]
    )

    return response.output_text


def txt_file(text):
    buffer = BytesIO()
    buffer.write(text.encode("utf-8"))
    buffer.seek(0)
    return buffer


def convert_json_to_excel(data):
    wb = Workbook()
    ws = wb.active
    ws.title = "Storyline"

    for variation in data["variations"]:

        ws.append([variation["story_variation"]])
        ws.append([f'Concept: {variation["concept"]}'])
        ws.append([f'Setting: {variation["setting"]}'])
        ws.append([f'Tone: {variation["tone"]}'])
        ws.append([""])

        headers = [
            "Shot",
            "Scene",
            "Visual",
            "Camera",
            "VO",
            "On-screen text",
            "Mood / Lighting",
            "Setting",
            "Outfit",
            "Duration"
        ]

        ws.append(headers)

        for shot in variation["shots"]:
            ws.append([
                shot["shot"],
                shot["scene"],
                shot["visual"],
                shot["camera"],
                shot["vo"],
                shot["on_screen_text"],
                shot["mood_lighting"],
                shot["setting"],
                shot["outfit"],
                shot["duration"]
            ])

        ws.append([])
        ws.append([])

    file_stream = BytesIO()
    wb.save(file_stream)
    file_stream.seek(0)
    return file_stream

def extract_ppt_text(file):
    prs = Presentation(file)
    slides_text = []
    for i, slide in enumerate(prs.slides):
        slide_text = []
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                slide_text.append(shape.text)
        slides_text.append(f"Slide {i+1}: " + " ".join(slide_text))
    return "\n".join(slides_text)


#ImageStyle
IMAGE_STYLE_OPTIONS = {
    "iPhone Front Camera": """
Image Style: Shot on an iPhone front camera with a smartphone video aesthetic. Wide-angle smartphone optics (approx. 24–26mm equivalent). Natural HDR processing with slightly lifted shadows and protected highlights. Subtle digital sharpening and smartphone noise reduction. Minor rolling exposure adjustments and small auto white-balance shifts. Light compression artifacts typical of mobile video.
Natural skin texture with mild computational smoothing, not beauty-filtered. Slight handheld micro-shake, imperfect stabilization. Clean but non-cinematic depth rendering (no dramatic bokeh). Casual, real-world phone video quality, not studio or film-grade.
No cinematic color grading, no film grain, no anamorphic lens, no DSLR or mirrorless look, no perfect stabilization, no studio lighting, no ultra-shallow depth of field.
""",

    "DSLR Camera": """
Image Style: Shot on a professional studio DSLR camera with a controlled, high-end photographic aesthetic. Standard prime or short zoom lens (approx. 35–50mm equivalent). Large-sensor imaging with high dynamic range, clean highlight roll-off, and deep color fidelity. Crisp optical sharpness with minimal digital processing and no computational artifacts.
Natural, detailed skin texture with realistic micro-contrast and no computational smoothing or beauty filtering. Precise manual white balance and consistent exposure. Stable framing with no handheld shake. Shallow but controlled depth of field with smooth, natural bokeh separation. Even, intentional studio lighting with soft key and controlled fill, producing clean shadows and dimensional subject separation.
No smartphone processing, no rolling exposure shifts, no compression artifacts, no handheld motion, no auto white-balance drift, no mobile HDR look, no casual or lo-fi aesthetic.
""",

    "iPhone Rear Camera": """
Image Style: Shot on an iPhone rear camera with a smartphone video aesthetic. Wide-angle rear camera optics (approx. 24–26mm equivalent). Natural HDR processing with slightly lifted shadows and protected highlights. Subtle digital sharpening and standard smartphone noise reduction. Minor rolling exposure adjustments and small auto white-balance shifts typical of rear camera capture. Light compression artifacts consistent with mobile video recording.
Natural skin texture with mild computational smoothing from the rear camera pipeline, not beauty-filtered. Slight handheld micro-shake with imperfect stabilization. Clean but non-cinematic depth rendering (no dramatic bokeh or portrait mode blur). Casual, real-world rear camera phone video quality, not studio or film-grade.
No cinematic color grading, no film grain, no anamorphic lens effects, no DSLR or mirrorless look, no perfect stabilization, no studio lighting, no ultra-shallow depth of field.
"""
}

#buat tab 1 output

def format_brand_brief_plain(data):

    p = data.get("product_brand_context", {})
    a = data.get("audience_market_context", {})
    c = data.get("campaign_distribution_context", {})

    lines = []

    lines.append("=== PRODUCT & BRAND CONTEXT ===\n")
    lines.append(f"Product Name: {p.get('product_name', 'unknown')}")
    lines.append(f"Product Category: {p.get('product_category', 'unknown')}")

    lines.append("\nProduct Value Propositions (Ranked):")
    for item in p.get("product_value_propositions_ranked", []):
        rank = item.get("rank", "unknown")
        value = item.get("value", "unknown")
        lines.append(f"  Rank {rank}: {value}")

    lines.append(f"\nProduct Usage: {p.get('product_usage', 'unknown')}")
    lines.append(f"Mandatory Brand Rules: {p.get('mandatory_brand_rules', 'unknown')}")

    lines.append("\n=== AUDIENCE & MARKET CONTEXT ===\n")
    lines.append(f"Target Market: {a.get('target_market', 'unknown')}")
    lines.append(f"Region: {a.get('region', 'unknown')}")
    lines.append(f"Language: {a.get('language', 'unknown')}")
    lines.append(f"Age & Persona: {a.get('age_persona', 'unknown')}")
    lines.append(f"Key Objections: {a.get('key_objections', 'unknown')}")
    lines.append(f"Primary Hook Angle: {a.get('primary_hook_angle', 'unknown')}")

    lines.append("\n=== CAMPAIGN & DISTRIBUTION CONTEXT ===\n")
    lines.append(f"Objective: {c.get('objective', 'unknown')}")
    lines.append(f"Platform: {c.get('platform', 'unknown')}")
    lines.append(f"Story Concept: {c.get('story_concept', 'unknown')}")
    lines.append(f"Tone: {c.get('tone', 'unknown')}")
    lines.append(f"CTA Intent: {c.get('cta_intent', 'unknown')}")

    return "\n".join(lines)

# ======================================
# UI
# ======================================

header_left, header_right = st.columns([6,1])

with header_left:
    st.image("header.png", width=220)

with header_right:

    st.markdown("<div style='height:30px'></div>", unsafe_allow_html=True)

    st.markdown(
        "<div style='display:flex; justify-content:flex-end;'>",
        unsafe_allow_html=True
    )

    stop_all = st.button("⏹", help="Stop / Reset")

    st.markdown("</div>", unsafe_allow_html=True)

    if stop_all:
        keys_to_clear = [
            "brand_brief_json",
            "brand_brief_plain",
            "storyline_versions",
            "revision_done",
            "character_prompt_output",
            "environment_prompt_output",
            "image_video_output"
        ]

        for k in keys_to_clear:
            st.session_state.pop(k, None)

        st.rerun()

if "active_tab" not in st.session_state:
    st.session_state["active_tab"] = 0

tab_index = st.session_state["active_tab"]

tab1, tab2, tab3 = st.tabs([
    "Brand Brief",
    "Storyline Generator",
    "Prompt Generator"
])



# ======================================
# TAB 1 — BRAND BRIEF
# ======================================

with tab1:
    controls, output = st.columns([1,2])

with controls:
    #st.header("Creative Spec Standardization")

    product_files = st.file_uploader(
    "Upload Product Assets (Images or PowerPoint)",
    type=["png","jpg","jpeg","ppt","pptx"],
    accept_multiple_files=True
    )

    product_desc = st.text_area("Product Description")
    audience = st.text_area("Audience")
    campaign = st.text_area("Campaign")

    if st.button("Generate Brand Brief"):

        with st.spinner("Generating Brand Brief..."):

            content = [
                {
                    "type": "input_text",
                    "text": f"""
            Convert the following into a standardized Brand Brief.

            Product Context:
            {product_desc}

            Audience Context:
            {audience}

            Campaign:
            {campaign}
            """
                }
            ]

            if product_files:

                for file in product_files:

                    if file.type in ["image/png", "image/jpeg"]:

                        image_bytes = file.read()
                        image_base64 = base64.b64encode(image_bytes).decode("utf-8")

                        content.append({
                            "type": "input_image",
                            "image_url": f"data:{file.type};base64,{image_base64}"
                        })

                    elif file.name.endswith(("ppt", "pptx")):

                        ppt_text = extract_ppt_text(file)

                        content.append({
                            "type": "input_text",
                            "text": f"\nProduct Deck Content:\n{ppt_text}"
                        })

            # Minimal, SDK-compatible call
            response = client.responses.create(
                model=MODEL_NAME,
                input=[
                    {"role": "system", "content": STANDARDIZATION_SYSTEM_PROMPT},
                    {"role": "user", "content": content}
                ]
            )

            result = response.output_text.strip()

            try:
                parsed = json.loads(result)
            except json.JSONDecodeError:
                # Auto-clean common markdown wrapping issues
                cleaned = result.replace("```json", "").replace("```", "").strip()
                try:
                    parsed = json.loads(cleaned)
                except json.JSONDecodeError:
                    st.error("Model did not return valid JSON.")
                    st.code(result)
                    st.stop()

            # Store in session state
            st.session_state["brand_brief_json"] = parsed
            st.session_state["brand_brief_plain"] = format_brand_brief_plain(parsed)
            
with output:
    # Display section
    if "brand_brief_json" in st.session_state:

        spec_tab1, spec_tab2 = st.tabs(["Plain Text", "JSON"])

        with spec_tab1:
            st.text_area(
                "Brand Brief (Readable)",
                st.session_state["brand_brief_plain"],
                height=500
            )

            st.download_button(
                label="Download Brand Brief (.txt)",
                data=txt_file(st.session_state["brand_brief_plain"]),
                file_name="brand_brief.txt",
                mime="text/plain"
            )

        with spec_tab2:
            st.json(st.session_state["brand_brief_json"])

# ======================================
# TAB 2 — STORYLINE GENERATOR (FINAL STABLE VERSIONED)
# ======================================

with tab2:
    controls, output = st.columns([1,2])
    
with controls:
    #st.header("Storyline Generator")

    st.markdown("""
        <div style="font-size:14px;font-weight:600;margin-bottom:4px;">
        Product Context
        </div>
        """, unsafe_allow_html=True)

    product_spec = st.text_area("Brand Brief (Copy the output from Tab 1 and paste it here)", height=150)

    product_image = st.file_uploader("Upload Product Image")

    st.markdown(
        "<hr style='margin-top:10px;margin-bottom:10px;border-color:#262730;'>",
        unsafe_allow_html=True
    )
    st.markdown("""
        <div style="font-size:14px;font-weight:600;margin-bottom:4px;">
        Story Settings
        </div>
        """, unsafe_allow_html=True)

    col1, col2, col3 = st.columns(3)

    with col1:
        num_stories = st.number_input("Stories", min_value=1, value=1)

    with col2:
        shots = st.number_input("Shots / Story", min_value=1, value=8)

    with col3:
        duration = st.selectbox("Duration", ["10 sec", "15 sec", "20 sec", "30 sec"])

    col4, col5 = st.columns(2)

    with col4:
        language = st.selectbox("Language", ["Indonesian","English"])

    with col5:
        text_language = st.selectbox("VO & Supers Language", ["Indonesian","English"])

    st.markdown(
        "<hr style='margin-top:10px;margin-bottom:10px;border-color:#262730;'>",
        unsafe_allow_html=True
    )

    st.markdown("""
    <div style="font-size:14px;font-weight:600;margin-bottom:4px;">
    Creative Direction
    </div>
    """, unsafe_allow_html=True)

    creative_direction = st.text_area("Tone, style, and narrative direction for the storyboard.", height=120)

    st.markdown(
        "<hr style='margin-top:10px;margin-bottom:10px;border-color:#262730;'>",
        unsafe_allow_html=True
    )

    st.markdown("""
    <div style="font-size:14px;font-weight:600;margin-bottom:4px;">
    Character Traits
    </div>
    """, unsafe_allow_html=True)

    character_traits = st.text_area("Traits defining speech, reactions, and behavior. Avoid visuals.", height=120)


    # creative_spec_input = st.text_area(
    #     "Product Creative Spec (Copy Paste Here)",
    #     value=st.session_state.get("creative_spec", ""),
    #     height=100
    # )

    # character_traits = st.text_area("Character Traits")
    # st.file_uploader("Upload Product Image", type=["png", "jpg", "jpeg"])

    # num_stories = st.number_input("Number of Stories", 1, 10, 1)
    # shots_per_story = st.number_input("Number of Shots Per Story", 1, 20, 8)

    # duration_each = st.selectbox(
    #     "Duration Each Story",
    #     options=["10 sec", "15 sec", "20 sec"],
    #     index=2
    # )

    # language = st.selectbox(
    #     "Language",
    #     options=["Indonesian", "English"],
    #     index=0
    # )

    # vo_language = st.selectbox(
    #     "VO & On-Screen Text Language",
    #     options=["Indonesian", "English"],
    #     index=0
    # )

    # other_info = st.text_area("Creative Direction")

    # ======================================
    # GENERATE STORYLINE
    # ======================================

    if st.button("Generate Storyline"):

        with st.spinner("Generating Storyline..."):
            
            user_text = f"""
            Product Knowledge:
            {product_spec}

            Character Traits:
            {character_traits}

            Creative Direction:
            {creative_direction}

            Requirements:
            - Number of variations: {num_stories}
            - Shots per variation: {shots}
            - Duration per story: {duration}
            - Language: {language}
            - VO Language: {text_language}
            """

            output = call_text_model(
                STORYLINE_SYSTEM_PROMPT,
                user_text,
            )

            try:
                parsed = json.loads(output)
            except json.JSONDecodeError:
                cleaned = output.replace("```json", "").replace("```", "").strip()
                try:
                    parsed = json.loads(cleaned)
                except:
                    st.error("Model did not return valid JSON.")
                    st.code(output)
                    st.stop()

            # Reset version history
            st.session_state["storyline_versions"] = [parsed]
            st.rerun()

    # ======================================
    # DISPLAY + DOWNLOAD
    # ======================================
with output:
    if "storyline_versions" in st.session_state:

        versions = st.session_state["storyline_versions"]

        # Always show latest version
        data = versions[-1]
        version_count = len(versions)

        with st.expander("Storyline JSON Output (Technical View)", expanded=False):
            st.json(data)

        st.subheader("Storyline Table Preview")

        for variation in data["variations"]:
            st.write(variation["story_variation"])
            st.write("Concept:", variation["concept"])
            st.write("Setting:", variation["setting"])
            st.write("Tone:", variation["tone"])

            df = pd.DataFrame(variation["shots"])
            st.dataframe(df, use_container_width=True)

        excel_file = convert_json_to_excel(data)

        # File naming logic
        if version_count == 1:
            file_name = "storyboard.xlsx"
        else:
            file_name = f"storyboard version {version_count - 1}.xlsx"

        st.download_button(
            label="Download Storyline (.xlsx)",
            data=excel_file,
            file_name=file_name,
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )

        # ======================================
        # REVISION SECTION
        # ======================================

        if st.session_state.get("revision_done"):
            st.success("Revision generated successfully.")
            del st.session_state["revision_done"]

        st.subheader("Revise Storyline")

        with st.form("revision_form"):

            revision_note = st.text_area(
                "Describe the revision clearly",
                placeholder="Example: Make tone more emotional but keep structure same."
            )

            submit_revision = st.form_submit_button("Apply Revision")

        if submit_revision:

            if not revision_note.strip():
                st.warning("Please enter a revision note.")
                st.stop()

            with st.spinner("Revising storyline..."):

                previous_json = json.dumps(data, indent=2)

                revision_prompt = f"""
You are editing an existing JSON object.

Current JSON:
{previous_json}

Apply this revision:
{revision_note}

Rules:
- Return FULL updated JSON.
- Keep structure identical.
- Do not remove fields.
- Do not change variation count.
- Do not change shot count.
- Return JSON only.
"""

                response = client.responses.create(
                    model=MODEL_NAME,
                    input=[
                        {"role": "system", "content": STORYLINE_SYSTEM_PROMPT},
                        {"role": "user", "content": revision_prompt}
                    ]
                )

                output = response.output_text.strip()

                try:
                    parsed = json.loads(output)
                except json.JSONDecodeError:
                    cleaned = output.replace("```json", "").replace("```", "").strip()
                    try:
                        parsed = json.loads(cleaned)
                    except:
                        st.error("Revision failed. Invalid JSON returned.")
                        st.code(output)
                        st.stop()

                # Append new version
                st.session_state["storyline_versions"].append(parsed)
                st.session_state["revision_done"] = True
                # Rerun to refresh UI
                st.rerun()

# ======================================
# TAB 3 — PROMPT GENERATOR
# ======================================

with tab3:
    controls, output = st.columns([1,2])

    with controls:

        #st.header("Prompt Generator")

        storyline_input = st.text_area(
            "Storyline (Copy the output from Tab 2 and paste it here)",
            height=200
        )

        brand_brief_input_pg = st.text_area(
            "Brand Brief",
            height=100
        )

        selected_style = st.selectbox(
            "Image Style",
            options=list(IMAGE_STYLE_OPTIONS.keys()),
            index=0
        )
        image_style_input = IMAGE_STYLE_OPTIONS[selected_style]

        ethnicity_outfit_input = st.text_area(
            "Character Design (Ethnicity, Outfit, Visual Identity)",
            placeholder="e.g. Javanese woman, pastel hijab, modern modest wear"
        )

        character_references = st.file_uploader(
            "Upload Character Reference Images",
            type=["png", "jpg", "jpeg"],
            accept_multiple_files=True
        )

        #st.caption("Limit 5MB per file • PNG, JPG, JPEG")
        # ======================================
        # SIDE-BY-SIDE BUTTONS HERE
        # ======================================

        col1, col2, col3 = st.columns(3)

        with col1:
            generate_char = st.button(
                "Character",
                use_container_width=True
            )

        with col2:
            generate_env = st.button(
                "Environment",
                use_container_width=True
            )

        with col3:
            generate_iv = st.button(
                "Image & Video",
                use_container_width=True
            )


        # ======================================
        # GENERATION LOGIC BELOW
        # ======================================

        if generate_char:
            st.session_state.pop("image_video_output", None)
            with st.spinner("Generating character prompts..."):

                character_prompt_input = f"""
        Storyline:
        {storyline_input}

        Brand Brief:
        {brand_brief_input_pg}

        Image Style:
        {image_style_input}

        Ethnicity & Outfit Notes:
        {ethnicity_outfit_input}
        """
                content = [
                    {"type": "input_text", "text": character_prompt_input}
                ]

                if character_references:
                    for file in character_references:

                        image_bytes = file.read()
                        image_base64 = base64.b64encode(image_bytes).decode("utf-8")

                        content.append({
                            "type": "input_image",
                            "image_url": f"data:image/jpeg;base64,{image_base64}"
                        })

                response = client.responses.create(
                    model=MODEL_NAME,
                    input=[
                        {"role": "system", "content": CHARACTER_SYSTEM_PROMPT},
                        {
                            "role": "user",
                            "content": content
                        }
                    ]
                )

                st.session_state["character_prompt_output"] = response.output_text

        if generate_iv:
            if "character_prompt_output" not in st.session_state:
                st.warning("Generate Character Prompts first.")
                st.stop()

            if "environment_prompt_output" not in st.session_state:
                st.warning("Generate Environment Prompts first.")
                st.stop()
            
            with st.spinner("Generating image & video prompts..."):

                iv_prompt_input = f"""
                Storyline:
                {storyline_input}

                Brand Brief:
                {brand_brief_input_pg}

                Character Design (LOCKED):
                {st.session_state.get("character_prompt_output","")}

                Environment Design (LOCKED):
                {st.session_state.get("environment_prompt_output","")}

                Image Style:
                {image_style_input}
                """
                response = client.responses.create(
                    model=MODEL_NAME,
                    input=[
                        {"role": "system", "content": IMAGE_VIDEO_SYSTEM_PROMPT},
                        {"role": "user", "content": iv_prompt_input}
                    ]
                )

                st.session_state["image_video_output"] = response.output_text

        if generate_env:

            # environment berubah → invalidate image/video saja
            st.session_state.pop("image_video_output", None)

            with st.spinner("Generating environment prompts..."):

                env_prompt_input = f"""
        Storyline:
        {storyline_input}

        Brand Brief:
        {brand_brief_input_pg}

        Image Style:
        {image_style_input}
        """

                response = client.responses.create(
                    model=MODEL_NAME,
                    input=[
                        {"role": "system", "content": ENVIRONMENT_SYSTEM_PROMPT},
                        {"role": "user", "content": env_prompt_input}
                    ]
                )

                st.session_state["environment_prompt_output"] = response.output_text


    with output:

        if (
            "character_prompt_output" in st.session_state
            or "image_video_output" in st.session_state
        ):

            #st.subheader("Generated Outputs")

            tab1, tab2, tab3 = st.tabs([
                "Character Prompts",
                "Environment Prompts",
                "Image & Video Prompts"
            ])

            with tab1:
                if "character_prompt_output" in st.session_state:

                    character_text = st.session_state["character_prompt_output"]

                    copy_button(character_text)

                    edited_character = st.text_area(
                        "",
                        value=character_text,
                        height=600,
                        label_visibility="collapsed"
                    )

                    # update if edited
                    if edited_character != character_text:
                        st.session_state["character_prompt_output"] = edited_character

                else:
                    st.info("No character prompts generated yet.")

            with tab2:
                if "environment_prompt_output" in st.session_state:

                    env_text = st.session_state["environment_prompt_output"]

                    copy_button(env_text)

                    edited_env = st.text_area(
                        "",
                        value=env_text,
                        height=600,
                        label_visibility="collapsed"
                    )

                    if edited_env != env_text:
                        st.session_state["environment_prompt_output"] = edited_env

                else:
                    st.info("No environment prompts generated yet.")

            with tab3:
                if "image_video_output" in st.session_state:

                    iv_text = st.session_state["image_video_output"]

                    copy_button(iv_text)

                    st.text_area(
                        "",
                        iv_text,
                        height=600,
                        label_visibility="collapsed"
                    )
                else:
                    st.info("No image & video prompts generated yet.")



    # # ======================================
    # # OUTPUT TABS
    # # ======================================

    # if (
    #     "character_prompt_output" in st.session_state
    #     or "image_video_output" in st.session_state
    # ):

    #     output_tab1, output_tab2 = st.tabs([
    #         "Character Prompts",
    #         "Image & Video Prompts Per Shot"
    #     ])
    #     # ------------------------------
    #     # CHARACTER TAB
    #     # ------------------------------
    #     with output_tab1:

    #         if "character_prompt_output" in st.session_state:
    #             character_text = st.session_state["character_prompt_output"]

    #             copy_button(character_text, "Copy Text")

    #             st.text_area(
    #                 "",
    #                 character_text,
    #                 height=600,
    #                 label_visibility="collapsed"
    #             )
    #         else:
    #             st.info("No character prompts generated yet.")


    #     # ------------------------------
    #     # IMAGE & VIDEO TAB
    #     # ------------------------------
    #     with output_tab2:

    #         if "image_video_output" in st.session_state:
    #             iv_text = st.session_state["image_video_output"]

    #             copy_button(iv_text, "Copy Text")

    #             st.text_area(
    #                 "",
    #                 iv_text,
    #                 height=600,
    #                 label_visibility="collapsed"
    #             )
    #         else:
    #             st.info("No image & video prompts generated yet.")


# ======================================
# STICKY FOOTER
# ======================================

st.markdown("""
<style>

/* prevent content from being hidden behind footer */
.block-container {
padding-bottom: 80px;
}

/* sticky footer */
.footer {
position: fixed;
bottom: 0;
left: 0;
width: 100%;
background-color: #0E1117;
text-align: center;
padding: 10px;
border-top: 1px solid #262730;
z-index: 1000;
}

.footer-line1 {
font-size: 12px;
color: #CBD5E1;
}

.footer-line2 {
font-size: 12px;
color: #CBD5E1;
font-style: italic;
margin-top: 2px;
}

</style>

<div class="footer">
    <div class="footer-line1">
        Version 0.3. Currently only support Custom AI Ads (Brand Brief, Storyline Creation, Character Prompt Generation,
        Image & Video Prompt Generation). Every frame is a delightful rebellion against the mundane. 
    </div>
</div>

""", unsafe_allow_html=True)